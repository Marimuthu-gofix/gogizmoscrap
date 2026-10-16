"""Buyback price dashboard (Flask).

Shows the latest Cashify and DOFY buyback price of every phone variant side
by side, with search, brand and comparison filters, Excel/CSV exports, the
history of checks, and buttons that run the scraper (buyback.py) in the
background.

Prices come from the local SQLite file, or from the remote price store
(hostinger_api/) when PRICE_API_URL and PRICE_API_KEY are set. The PHP
dashboard in hostinger_site/ is a port of this file and its template; keep
the two in step.
"""
from __future__ import annotations

import csv
import io
import json
import os
import re
import sqlite3
import subprocess
import sys
import threading
import uuid
from collections import defaultdict
from contextlib import contextmanager
from datetime import datetime, timezone
from functools import lru_cache
from pathlib import Path
from typing import Any, Iterator
from zoneinfo import ZoneInfo

from flask import Flask, Response, jsonify, redirect, render_template, request, url_for
from markupsafe import Markup

from buyback import SIZE_RE, TABLE, canonical_brand, connect_db
from price_store import ROW_FIELDS, RemotePriceStore

ROOT = Path(__file__).resolve().parent
DEFAULT_DB = ROOT / "buyback_history.sqlite3"
HOSTED_READ_ONLY = bool(os.environ.get("VERCEL"))
# When PRICE_API_URL / PRICE_API_KEY are set, prices are read from the remote
# store (hostinger_api/) instead of the local SQLite file.
REMOTE_STORE = RemotePriceStore.from_env()
DISPLAY_TIMEZONE = ZoneInfo("Asia/Kolkata")
PLATFORMS = ("Cashify", "DOFY")
SCRAPE_PLATFORMS = {"all": "Cashify + DOFY", "cashify": "Cashify", "dofy": "DOFY"}
SCRAPE_BRANDS = [
    "Apple", "Samsung", "Google", "OnePlus", "Xiaomi", "POCO", "OPPO", "vivo", "realme", "Motorola", "iQOO",
    "Nothing", "CMF", "Infinix", "Tecno", "Nokia", "Honor", "Asus", "Lava", "Huawei", "LG", "Nubia", "Lenovo",
]
PER_PAGE_OPTIONS = [25, 50, 100, 250]
HISTORY_PER_PAGE_OPTIONS = [10, 25, 50, 100]
HISTORY_PER_PAGE_DEFAULT = 10
COMPARE_OPTIONS = [
    ("", "All variants"),
    ("both", "On both platforms"),
    ("cashify_higher", "Cashify pays more"),
    ("dofy_higher", "DOFY pays more"),
    ("cashify_only", "Only on Cashify"),
    ("dofy_only", "Only on DOFY"),
]
COMPARE_LABELS = dict(COMPARE_OPTIONS)
SORT_KEYS = ("brand", "model", "variant", "cashify", "dofy", "gap", "date")
# Leading words that only repeat the brand ("Xiaomi Redmi 10" vs "Redmi 10").
MODEL_PREFIX_WORDS = {
    "apple", "samsung", "google", "xiaomi", "oppo", "vivo", "realme", "oneplus", "motorola", "iqoo", "nothing",
    "poco", "infinix", "tecno", "nokia", "lava", "honor", "huawei", "lg", "nubia", "asus", "lenovo", "cmf",
}
NETWORK_WORDS = {"5g", "4g", "lte"}

app = Flask(__name__)
job_lock = threading.Lock()
job_state: dict[str, Any] = {
    "id": None,
    "running": False,
    "started_at": None,
    "finished_at": None,
    "returncode": None,
    "message": "Ready",
    "output": "",
    "progress": {"done": 0, "rows": 0, "errors": 0, "current": "", "scope": ""},
}


# ------------------------------------------------------------- formatting
def money(value: float | None, currency: str | None = "INR") -> str:
    if value is None:
        return "-"
    prefix = "Rs " if currency in (None, "INR") else f"{currency} "
    return f"{prefix}{round(value):,.0f}"


@app.template_filter("money")
def money_filter(value: float | None, currency: str | None = "INR") -> str:
    return money(value, currency)


ICONS = {
    "refresh": '<path d="M3 12a9 9 0 0 1 15.3-6.4M21 12a9 9 0 0 1-15.3 6.4"/><path d="M3 3v5h5M21 21v-5h-5"/>',
    "download": '<path d="M12 3v12m0 0-4-4m4 4 4-4"/><path d="M4 19h16"/>',
    "dot": '<circle cx="12" cy="12" r="5"/>',
    "play": '<path d="M6 4l14 8-14 8V4z"/>',
    "search": '<circle cx="11" cy="11" r="7"/><path d="m21 21-4.3-4.3"/>',
    "search-sm": '<circle cx="11" cy="11" r="7"/><path d="m21 21-4.3-4.3"/>',
    "target": '<circle cx="12" cy="12" r="8"/><circle cx="12" cy="12" r="4"/><circle cx="12" cy="12" r="0.5"/>',
    "box": '<path d="M21 8 12 3 3 8v8l9 5 9-5V8Z"/><path d="M3 8l9 5 9-5M12 13v8"/>',
    "layers": '<path d="m12 2 9 5-9 5-9-5 9-5Z"/><path d="m3 12 9 5 9-5M3 17l9 5 9-5"/>',
    "clock": '<circle cx="12" cy="12" r="9"/><path d="M12 7v5l3 3"/>',
    "tag": '<path d="M20 12 12.5 19.5a2 2 0 0 1-2.8 0l-6.2-6.2a2 2 0 0 1 0-2.8L11 3h9v9Z"/><circle cx="15" cy="8" r="1.5"/>',
    "scale": '<path d="M12 3v18M7 21h10M5 7h14"/><path d="m5 7-3 7a3 3 0 0 0 6 0L5 7ZM19 7l-3 7a3 3 0 0 0 6 0l-3-7Z"/>',
    "chevron": '<path d="m9 6 6 6-6 6"/>',
    "empty": '<path d="M4 7h16l-1.5 12.5a2 2 0 0 1-2 1.8H7.5a2 2 0 0 1-2-1.8L4 7Z"/><path d="M9 7V5a3 3 0 0 1 6 0v2"/>',
    "check": '<path d="m5 12 5 5 9-10"/>',
    "alert": '<path d="M12 3 2 21h20L12 3Z"/><path d="M12 10v4M12 17h.01"/>',
    "asc": '<path d="m6 15 6-6 6 6"/>',
    "desc": '<path d="m6 9 6 6 6-6"/>',
    "sort-none": '<path d="m8 9 4-4 4 4M8 15l4 4 4-4" opacity="0.4"/>',
    "user": '<circle cx="12" cy="8" r="4"/><path d="M4 21a8 8 0 0 1 16 0"/>',
}


@app.template_global("icon")
def icon(name: str) -> Markup:
    body = ICONS.get(name, ICONS["dot"])
    return Markup(
        f'<svg class="icon icon-{name}" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" '
        f'stroke-linecap="round" stroke-linejoin="round" aria-hidden="true">{body}</svg>'
    )


@app.route("/favicon.ico")
def favicon() -> Response:
    return app.send_static_file("favicon.ico")


@app.template_global()
def static_version(filename: str) -> int:
    try:
        return int((ROOT / "static" / filename).stat().st_mtime)
    except OSError:
        return 0


def parse_iso(value: str | None) -> datetime | None:
    try:
        parsed = datetime.fromisoformat((value or "").replace("Z", "+00:00"))
    except ValueError:
        return None
    return parsed if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)


@app.template_filter("localdate")
def localdate_filter(value: str | None) -> str:
    """"28 Sep 2026, 11:17 AM" in India time."""
    parsed = parse_iso(value)
    return parsed.astimezone(DISPLAY_TIMEZONE).strftime("%d %b %Y, %I:%M %p") if parsed else (value or "")


@app.template_filter("localdate_short")
def localdate_short_filter(value: str | None) -> str:
    parsed = parse_iso(value)
    return parsed.astimezone(DISPLAY_TIMEZONE).strftime("%d %b, %I:%M %p") if parsed else (value or "")


def local_day(value: str | datetime | None) -> str:
    parsed = value if isinstance(value, datetime) else parse_iso(value)
    return parsed.astimezone(DISPLAY_TIMEZONE).strftime("%Y-%m-%d") if parsed else ""


# ----------------------------------------------------------- name matching
def clean_text(value: str | None) -> str:
    return re.sub(r"\s+", " ", (value or "").replace("\xa0", " ")).strip().lower()


def clean_label(value: str | None) -> str:
    return re.sub(r"\s+", " ", (value or "").replace("\xa0", " ")).strip()


def brand_key(brand: str | None) -> str:
    return canonical_brand(brand).lower()


@lru_cache(maxsize=20000)
def model_key(model: str | None, loose: bool = False) -> str:
    """"Samsung Galaxy S24+ 5G" -> "galaxy s24 plus 5g" (loose: "galaxy s24 plus").

    The exact key keeps "5G" because some brands sell a 4G and a 5G phone
    under one name ("Vivo T1" and "Vivo T1 5G"); the loose key is only used
    when it points at a single model on each platform.
    """
    text = clean_text(model).replace("+", " plus ")
    words = re.sub(r"[^a-z0-9 ]+", " ", text).split()
    while len(words) > 1 and words[0] in MODEL_PREFIX_WORDS:
        words.pop(0)
    if loose:
        words = [word for word in words if word not in NETWORK_WORDS] or words
    return " ".join(words)


def variant_sizes(variant: str | None) -> tuple[float | None, float | None]:
    """(RAM GB, storage GB) from "8 GB/256 GB"; a single size is storage."""
    sizes = []
    for number, unit in SIZE_RE.findall(variant or ""):
        factor = {"TB": 1024.0, "GB": 1.0, "MB": 1 / 1024}[unit.upper()]
        sizes.append(float(number) * factor)
    if not sizes:
        return None, None
    if len(sizes) == 1:
        return None, sizes[0]
    return sizes[0], sizes[-1]


def matches_search(entry: dict[str, Any], search: str) -> bool:
    """Every word of the search must appear somewhere in the row."""
    words = clean_text(search).split()
    if not words:
        return True
    return all(word in entry["search_text"] for word in words)


# ---------------------------------------------------------------- loading
def db_path() -> Path:
    return Path(app.config.get("DB_PATH") or os.environ.get("DB_PATH") or DEFAULT_DB)


def data_source_label() -> str:
    return REMOTE_STORE.base_url if REMOTE_STORE is not None else str(db_path().name)


@contextmanager
def local_conn() -> Iterator[sqlite3.Connection]:
    conn = connect_db(db_path())
    conn.row_factory = sqlite3.Row
    try:
        yield conn
    finally:
        conn.close()


LATEST_SQL = f"""
    SELECT p.id, p.{', p.'.join(ROW_FIELDS)}
    FROM {TABLE} p
    JOIN (
        SELECT platform, brand, model, variant, MAX(date) AS date
        FROM {TABLE} WHERE status = 'ok'
        GROUP BY platform, brand, model, variant
    ) m ON m.platform = p.platform AND m.brand = p.brand AND m.model = p.model AND m.variant = p.variant AND m.date = p.date
    WHERE p.status = 'ok'
"""


def store_status() -> dict[str, Any]:
    if REMOTE_STORE is not None:
        return REMOTE_STORE.status()
    with local_conn() as conn:
        row = conn.execute(f"SELECT COUNT(*) AS count, MAX(id) AS max_id, MAX(date) AS max_date FROM {TABLE}").fetchone()
    return {"count": int(row["count"]), "max_id": int(row["max_id"] or 0), "max_date": row["max_date"]}


def fetch_latest_rows() -> list[dict[str, Any]]:
    """The newest successful price of every (platform, brand, model, variant)."""
    if REMOTE_STORE is not None:
        return REMOTE_STORE.latest()
    with local_conn() as conn:
        return [dict(row) for row in conn.execute(LATEST_SQL)]


def fetch_checks(filters: dict[str, Any], offset: int, limit: int) -> tuple[list[dict[str, Any]], int]:
    """One page of the check history (newest first), filtered like the page."""
    if REMOTE_STORE is not None:
        return REMOTE_STORE.checks(
            offset=offset, limit=limit, status=filters["status"], q=filters["search"],
            brands=filters["brands"], platform=filters["platform"],
        )
    where, params = checks_where(filters)
    with local_conn() as conn:
        total = int(conn.execute(f"SELECT COUNT(*) FROM {TABLE}{where}", params).fetchone()[0])
        rows = conn.execute(
            f"SELECT {', '.join(ROW_FIELDS)} FROM {TABLE}{where} ORDER BY date DESC, id DESC LIMIT ? OFFSET ?",
            [*params, limit, offset],
        ).fetchall()
    return [dict(row) for row in rows], total


def checks_where(filters: dict[str, Any]) -> tuple[str, list[Any]]:
    """The WHERE clause for the history. hostinger_api/index.php builds the
    same one for the remote store."""
    where: list[str] = []
    params: list[Any] = []
    if filters["brands"]:
        where.append(f"LOWER(brand) IN ({', '.join('?' for _ in filters['brands'])})")
        params.extend(brand.lower() for brand in filters["brands"])
    if filters["status"] in ("ok", "error"):
        where.append("status = ?")
        params.append(filters["status"])
    if filters["platform"] in PLATFORMS:
        where.append("platform = ?")
        params.append(filters["platform"])
    for word in clean_text(filters["search"]).split()[:6]:
        # "!" escapes LIKE wildcards the same way in SQLite and MySQL.
        like = "%" + word.replace("!", "!!").replace("%", "!%").replace("_", "!_") + "%"
        where.append(
            "(LOWER(brand) LIKE ? ESCAPE '!' OR LOWER(model) LIKE ? ESCAPE '!' OR LOWER(variant) LIKE ? ESCAPE '!' "
            "OR LOWER(platform) LIKE ? ESCAPE '!')"
        )
        params.extend([like] * 4)
    return (" WHERE " + " AND ".join(where)) if where else "", params


# ------------------------------------------------------------ comparison
def compare_rows(latest: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Pair every Cashify variant with the same phone and storage on DOFY.

    Models are paired by brand and exact model key first, then by the loose
    key ("5G" dropped) when it is unambiguous on both sides. Inside a model
    pair, variants are paired by storage, and by RAM when both sides state it.
    Anything left over is listed for its one platform.
    """
    models: dict[str, dict[tuple[str, str], list[dict[str, Any]]]] = {name: defaultdict(list) for name in PLATFORMS}
    for row in latest:
        if row.get("platform") not in models or row.get("price_value") is None:
            continue
        row = dict(row)
        row["price_value"] = float(row["price_value"])
        row["ram"], row["storage"] = variant_sizes(row.get("variant"))
        models[row["platform"]][(brand_key(row["brand"]), model_key(row["model"]))].append(row)

    cashify, dofy = models["Cashify"], models["DOFY"]
    pairs: list[tuple[tuple[str, str] | None, tuple[str, str] | None]] = []
    open_c, open_d = set(cashify), set(dofy)
    for key in sorted(open_c & open_d):
        pairs.append((key, key))
    open_c -= set(dofy)
    open_d -= set(cashify)
    loose_c: dict[tuple[str, str], list[tuple[str, str]]] = defaultdict(list)
    loose_d: dict[tuple[str, str], list[tuple[str, str]]] = defaultdict(list)
    for key in open_c:
        loose_c[(key[0], model_key(key[1], loose=True))].append(key)
    for key in open_d:
        loose_d[(key[0], model_key(key[1], loose=True))].append(key)
    for loose, keys in loose_c.items():
        others = loose_d.get(loose) or []
        if len(keys) == 1 and len(others) == 1:
            pairs.append((keys[0], others[0]))
            open_c.discard(keys[0])
            open_d.discard(others[0])
    pairs.extend((key, None) for key in sorted(open_c))
    pairs.extend((None, key) for key in sorted(open_d))

    entries = []
    for c_key, d_key in pairs:
        c_rows = sorted(cashify.get(c_key, []) if c_key else [], key=lambda r: (r["storage"] or 0, r["ram"] or 0))
        d_rows = sorted(dofy.get(d_key, []) if d_key else [], key=lambda r: (r["storage"] or 0, r["ram"] or 0))
        group = f"{(c_key or d_key)[0]}|{(c_key or d_key)[1]}"
        used: set[int] = set()
        for c_row in c_rows:
            match = None
            for index, d_row in enumerate(d_rows):
                if index in used or d_row["storage"] != c_row["storage"]:
                    continue
                if c_row["ram"] is not None and d_row["ram"] is not None and c_row["ram"] != d_row["ram"]:
                    continue
                if match is None or (d_row["ram"] == c_row["ram"] and d_rows[match]["ram"] != c_row["ram"]):
                    match = index
            if match is not None:
                used.add(match)
            entries.append(make_entry(group, c_row, d_rows[match] if match is not None else None))
        for index, d_row in enumerate(d_rows):
            if index not in used:
                entries.append(make_entry(group, None, d_row))
    return entries


def make_entry(group: str, cashify: dict[str, Any] | None, dofy: dict[str, Any] | None) -> dict[str, Any]:
    first = cashify or dofy
    assert first is not None
    variants = [row["variant"] for row in (cashify, dofy) if row and row.get("variant")]
    # The more specific name wins: "4 GB/128 GB" over "128 GB".
    variant = max(variants, key=lambda v: (len(SIZE_RE.findall(v)), len(v))) if variants else ""
    c_value = cashify["price_value"] if cashify else None
    d_value = dofy["price_value"] if dofy else None
    best, gap, gap_pct = None, None, None
    if c_value is not None and d_value is not None:
        gap = abs(c_value - d_value)
        best = "Same" if gap == 0 else ("Cashify" if c_value > d_value else "DOFY")
        low = min(c_value, d_value)
        gap_pct = round(100 * gap / low, 1) if low else None
    dates = [row["date"] for row in (cashify, dofy) if row]
    entry = {
        "group": group,
        "brand": canonical_brand(first["brand"]),
        "brand_key": brand_key(first["brand"]),
        "model": clean_label((cashify or dofy)["model"]),
        "model_key": group.split("|", 1)[1],
        "variant": variant,
        "storage": first["storage"],
        "ram": next((row["ram"] for row in (cashify, dofy) if row and row["ram"] is not None), None),
        "cashify": cashify,
        "dofy": dofy,
        "cashify_value": c_value,
        "dofy_value": d_value,
        "best": best,
        "gap": gap,
        "gap_pct": gap_pct,
        "best_value": max(v for v in (c_value, d_value) if v is not None),
        "date": max(dates),
        "platforms": "both" if cashify and dofy else ("cashify" if cashify else "dofy"),
    }
    entry["day"] = local_day(entry["date"])
    names = " ".join(row["model"] for row in (cashify, dofy) if row)
    entry["search_text"] = clean_text(f"{entry['brand']} {names} {variant} {' '.join(v for v in variants)}")
    return entry


def compare_filter(entry: dict[str, Any], compare: str) -> bool:
    if compare == "both":
        return entry["platforms"] == "both"
    if compare == "cashify_higher":
        return entry["best"] == "Cashify"
    if compare == "dofy_higher":
        return entry["best"] == "DOFY"
    if compare == "cashify_only":
        return entry["platforms"] == "cashify"
    if compare == "dofy_only":
        return entry["platforms"] == "dofy"
    return True


def last_update_summary(latest: list[dict[str, Any]], gap_minutes: int = 15) -> dict[str, Any]:
    """How many prices the most recent scrape refreshed, and when. Counting
    back from the newest price, rows belong to the same scrape while
    consecutive check times are less than `gap_minutes` apart."""
    stamps = sorted((parsed for parsed in (parse_iso(row.get("date")) for row in latest) if parsed), reverse=True)
    if not stamps:
        return {"count": 0, "date": None}
    count = 1
    for newer, older in zip(stamps, stamps[1:]):
        if (newer - older).total_seconds() > gap_minutes * 60:
            break
        count += 1
    return {"count": count, "date": stamps[0].isoformat(timespec="seconds")}


_PRICE_DATA_LOCK = threading.Lock()
_PRICE_DATA: dict[str, Any] = {"stamp": None}


def load_price_data() -> dict[str, Any]:
    """Comparison rows plus the picker lists, rebuilt only when the store
    changes (row count or newest id), so most page loads are one status call."""
    status = store_status()
    stamp = f"{status.get('count')}|{status.get('max_id')}"
    with _PRICE_DATA_LOCK:
        if _PRICE_DATA.get("stamp") == stamp:
            return _PRICE_DATA
    latest = fetch_latest_rows()
    entries = compare_rows(latest)
    model_options = sorted({entry["model"] for entry in entries}, key=str.lower)
    brands = sorted({*SCRAPE_BRANDS, *(entry["brand"] for entry in entries)}, key=str.lower)
    data = {
        "stamp": stamp,
        "entries": entries,
        "total_rows": int(status.get("count") or 0),
        "last_update": last_update_summary(latest),
        "brands": brands,
        "model_options": model_options,
        "suggestions": model_options[:800],
    }
    with _PRICE_DATA_LOCK:
        _PRICE_DATA.clear()
        _PRICE_DATA.update(data)
    return data


def sort_value(entry: dict[str, Any], sort: str) -> Any:
    missing = float("-inf")
    if sort == "brand":
        return (entry["brand_key"], clean_text(entry["model"]), entry["storage"] or 0)
    if sort == "variant":
        return (entry["storage"] or 0, entry["ram"] or 0)
    if sort == "cashify":
        return entry["cashify_value"] if entry["cashify_value"] is not None else missing
    if sort == "dofy":
        return entry["dofy_value"] if entry["dofy_value"] is not None else missing
    if sort == "gap":
        return entry["gap"] if entry["gap"] is not None else missing
    if sort == "date":
        return entry["date"] or ""
    return (clean_text(entry["model"]), entry["brand_key"], entry["storage"] or 0, entry["ram"] or 0)


def group_by_model(entries: list[dict[str, Any]]) -> list[tuple[dict[str, Any], bool, int, str, str]]:
    """(entry, group_start, group_size, group_id, coverage) for consecutive
    rows of one model, so the page can fold a model's variants together."""
    grouped = []
    index = 0
    group_number = 0
    while index < len(entries):
        end = index
        while end < len(entries) and entries[end]["group"] == entries[index]["group"]:
            end += 1
        size = end - index
        platforms = {entry["platforms"] for entry in entries[index:end]}
        coverage = "Cashify + DOFY" if "both" in platforms or platforms == {"cashify", "dofy"} else (
            "Cashify only" if platforms == {"cashify"} else "DOFY only"
        )
        for position in range(index, end):
            grouped.append((entries[position], position == index, size, f"g{group_number}", coverage))
        group_number += 1
        index = end
    return grouped


def page_info(total: int, page: int, per_page: int, window: int = 2) -> dict[str, Any]:
    total_pages = max(1, (total + per_page - 1) // per_page)
    page = min(max(page, 1), total_pages)
    start = (page - 1) * per_page
    end = min(start + per_page, total)
    numbers = sorted({1, total_pages, *(p for p in range(page - window, page + window + 1) if 1 <= p <= total_pages)})
    return {
        "page": page, "per_page": per_page, "total": total, "total_pages": total_pages,
        "slice_start": start, "slice_end": end, "start": start + 1 if total else 0, "end": end,
        "has_prev": page > 1, "has_next": page < total_pages,
        "prev_page": max(1, page - 1), "next_page": min(total_pages, page + 1), "page_numbers": numbers,
    }


def load_dashboard(
    search: str = "",
    selected_brands: list[str] | None = None,
    compare: str = "",
    page: int = 1,
    per_page: int = 100,
    status: str = "",
    platform: str = "",
    sort: str = "model",
    sort_dir: str = "asc",
    history_page: int = 1,
    history_per_page: int = HISTORY_PER_PAGE_DEFAULT,
    paginate: bool = True,
    with_history: bool = True,
) -> dict[str, Any]:
    data = load_price_data()
    all_entries = data["entries"]
    selected_brands = selected_brands or []
    wanted = {brand.lower() for brand in selected_brands}
    entries = [
        entry for entry in all_entries
        if (not wanted or entry["brand"].lower() in wanted)
        and compare_filter(entry, compare)
        and matches_search(entry, search)
    ]

    sort_key = sort if sort in SORT_KEYS else "model"
    reverse = sort_dir == "desc"
    entries.sort(key=lambda entry: sort_value(entry, sort_key), reverse=reverse)
    if sort_key in ("model", "brand", "date"):
        # Keep every model's variants together, smallest storage first, in
        # the order the models first appear.
        first_seen: dict[str, int] = {}
        for entry in entries:
            first_seen.setdefault(entry["group"], len(first_seen))
        entries.sort(key=lambda entry: (first_seen[entry["group"]], entry["storage"] or 0, entry["ram"] or 0))

    total = len(entries)
    per_page = per_page if per_page in PER_PAGE_OPTIONS else 100
    if not paginate:
        per_page, page = max(total, 1), 1
    pagination = page_info(total, page, per_page)
    paged = group_by_model(entries[pagination["slice_start"]:pagination["slice_end"]])

    history: list[dict[str, Any]] = []
    history_pagination = page_info(0, 1, HISTORY_PER_PAGE_DEFAULT)
    if with_history:
        history_per_page = history_per_page if history_per_page in HISTORY_PER_PAGE_OPTIONS else HISTORY_PER_PAGE_DEFAULT
        filters = {"brands": selected_brands, "status": status, "search": search, "platform": platform}
        history, total_checks = fetch_checks(filters, (max(history_page, 1) - 1) * history_per_page, history_per_page)
        history_pagination = page_info(total_checks, history_page, history_per_page)
        if not history and total_checks:
            # The requested page is past the end; show the last one instead.
            history, total_checks = fetch_checks(filters, history_pagination["slice_start"], history_per_page)

    shared = [entry for entry in entries if entry["platforms"] == "both"]
    today = local_day(datetime.now(timezone.utc))
    return {
        "latest": paged,
        "history": history,
        "history_pagination": history_pagination,
        "brands": data["brands"],
        "per_page_options": PER_PAGE_OPTIONS,
        "history_per_page_options": HISTORY_PER_PAGE_OPTIONS,
        "compare_options": COMPARE_OPTIONS,
        "platform_options": PLATFORMS,
        "scrape_platforms": SCRAPE_PLATFORMS,
        "scrape_brands": SCRAPE_BRANDS,
        "suggestions": data["suggestions"],
        "sort": sort_key,
        "sort_dir": "desc" if reverse else "asc",
        "selected_status": status,
        "selected_platform": platform,
        "selected_brands": selected_brands,
        "selected_compare": compare,
        "pagination": pagination,
        "stats": {
            "tracked": total,
            "models": len({entry["group"] for entry in entries}),
            "brands": len({entry["brand"] for entry in entries}),
            "both": len(shared),
            "cashify_higher": sum(1 for entry in shared if entry["best"] == "Cashify"),
            "dofy_higher": sum(1 for entry in shared if entry["best"] == "DOFY"),
            "updated_today": sum(1 for entry in all_entries if entry["day"] == today),
            "last_update": data["last_update"],
        },
        "data_source": data_source_label(),
        "remote_error": None,
        "hosted_read_only": HOSTED_READ_ONLY,
        "job": current_job_state(),
    }


# ------------------------------------------------------------------ filters
def parse_positive_int(value: str | None, default: int) -> int:
    try:
        parsed = int(value or "")
    except ValueError:
        return default
    return parsed if parsed > 0 else default


def parse_filters() -> dict[str, Any]:
    brands: list[str] = []
    for value in request.args.getlist("brand") + request.args.getlist("brand[]"):
        cleaned = value.strip()
        if cleaned and cleaned not in brands:
            brands.append(cleaned)
    compare = request.args.get("compare", "").strip()
    platform = request.args.get("platform", "").strip()
    return {
        "search": request.args.get("q", "").strip(),
        "brands": brands,
        "compare": compare if compare in COMPARE_LABELS else "",
        "status": request.args.get("status", "").strip(),
        "platform": platform if platform in PLATFORMS else "",
        "page": parse_positive_int(request.args.get("page"), 1),
        "per_page": parse_positive_int(request.args.get("per_page"), 100),
        "history_page": parse_positive_int(request.args.get("hpage"), 1),
        "history_per_page": parse_positive_int(request.args.get("hper_page"), HISTORY_PER_PAGE_DEFAULT),
        "sort": request.args.get("sort", "model").strip() or "model",
        "sort_dir": request.args.get("dir", "asc").strip() or "asc",
    }


def build_query(filters: dict[str, Any], **overrides: Any) -> dict[str, Any]:
    query = {
        "q": filters["search"], "brand": filters["brands"], "compare": filters["compare"],
        "status": filters["status"], "platform": filters["platform"], "per_page": filters["per_page"],
        "hper_page": filters["history_per_page"], "hpage": filters["history_page"],
        "sort": filters["sort"], "dir": filters["sort_dir"],
    }
    query.update(overrides)
    return {key: value for key, value in query.items() if value not in ("", None, [])}


def build_active_filters(filters: dict[str, Any]) -> list[dict[str, str]]:
    active = []
    if filters["search"]:
        active.append({"label": f'Search "{filters["search"]}"', "remove_url": url_for("index", **build_query(filters, q="", page=1))})
    for selected in filters["brands"]:
        rest = [brand for brand in filters["brands"] if brand != selected]
        active.append({"label": selected, "remove_url": url_for("index", **build_query(filters, brand=rest, page=1))})
    if filters["compare"]:
        active.append({"label": COMPARE_LABELS[filters["compare"]], "remove_url": url_for("index", **build_query(filters, compare="", page=1))})
    if filters["platform"]:
        active.append({"label": f"Checks: {filters['platform']}", "remove_url": url_for("index", **build_query(filters, platform="", page=1))})
    if filters["status"] in ("ok", "error"):
        label = "Checks: available" if filters["status"] == "ok" else "Checks: errors"
        active.append({"label": label, "remove_url": url_for("index", **build_query(filters, status="", page=1))})
    return active


def sort_headers(filters: dict[str, Any]) -> list[dict[str, Any]]:
    headers = []
    for key, label in (("brand", "Brand"), ("model", "Model"), ("variant", "Variant"), ("cashify", "Cashify"),
                       ("dofy", "DOFY"), ("gap", "Best Offer"), ("date", "Last Checked")):
        active = filters["sort"] == key
        next_dir = "asc" if active and filters["sort_dir"] == "desc" else "desc"
        if not active and key in ("brand", "model", "variant"):
            next_dir = "asc"
        headers.append({
            "key": key, "label": label, "active": active,
            "url": url_for("index", **build_query(filters, sort=key, dir=next_dir, page=1)),
            "icon": ("asc" if filters["sort_dir"] == "asc" else "desc") if active else "sort-none",
        })
    return headers


# ------------------------------------------------------------------ routes
@app.route("/")
def index() -> str:
    filters = parse_filters()
    try:
        data = load_dashboard(
            search=filters["search"], selected_brands=filters["brands"], compare=filters["compare"],
            page=filters["page"], per_page=filters["per_page"], status=filters["status"], platform=filters["platform"],
            sort=filters["sort"], sort_dir=filters["sort_dir"],
            history_page=filters["history_page"], history_per_page=filters["history_per_page"],
        )
    except Exception as exc:
        if REMOTE_STORE is None:
            raise
        app.logger.exception("Could not read the remote price store")
        data = empty_dashboard(f"{type(exc).__name__}: {exc}")
    history_query = build_query(filters, page=data["pagination"]["page"])
    history_query.pop("hpage", None)
    return render_template(
        "dashboard.html",
        search=filters["search"],
        selected_per_page=data["pagination"]["per_page"],
        active_filters=build_active_filters(filters),
        sort_headers=sort_headers(filters),
        base_query=build_query(filters),
        history_query=history_query,
        **data,
    )


def empty_dashboard(error: str) -> dict[str, Any]:
    """What the page shows when the remote store cannot be reached."""
    empty = page_info(0, 1, 100)
    return {
        "latest": [], "history": [], "history_pagination": page_info(0, 1, HISTORY_PER_PAGE_DEFAULT),
        "brands": SCRAPE_BRANDS, "per_page_options": PER_PAGE_OPTIONS, "history_per_page_options": HISTORY_PER_PAGE_OPTIONS,
        "compare_options": COMPARE_OPTIONS, "platform_options": PLATFORMS, "scrape_platforms": SCRAPE_PLATFORMS,
        "scrape_brands": SCRAPE_BRANDS, "suggestions": [], "sort": "model", "sort_dir": "asc", "selected_status": "",
        "selected_platform": "", "selected_brands": [], "selected_compare": "", "pagination": empty,
        "stats": {"tracked": 0, "models": 0, "brands": 0, "both": 0, "cashify_higher": 0, "dofy_higher": 0,
                  "updated_today": 0, "last_update": {"count": 0, "date": None}},
        "data_source": data_source_label(), "remote_error": error, "hosted_read_only": HOSTED_READ_ONLY,
        "job": current_job_state(),
    }


def export_rows(scope: str, model: str = "") -> tuple[list[dict[str, Any]], str]:
    """Comparison rows for an export plus a short name for the file.

    scope "all"   -> every tracked variant
    scope "view"  -> whatever the dashboard filters currently show
    scope "model" -> one phone, every variant
    """
    filters = parse_filters()
    if scope == "view":
        data = load_dashboard(
            search=filters["search"], selected_brands=filters["brands"], compare=filters["compare"],
            sort="model", sort_dir="asc", paginate=False, with_history=False,
        )
        return [item[0] for item in data["latest"]], "current_view"
    data = load_dashboard(sort="model", sort_dir="asc", paginate=False, with_history=False)
    rows = [item[0] for item in data["latest"]]
    if scope != "model" or not model.strip():
        return rows, "all_models"
    wanted = model_key(model)
    chosen = [row for row in rows if row["model_key"] == wanted]
    if not chosen:
        chosen = [row for row in rows if model_key(row["model"], loose=True) == model_key(model, loose=True)]
    if not chosen:
        chosen = [row for row in rows if matches_search(row, model)]
    label = re.sub(r"[^a-z0-9]+", "_", clean_label(model).lower()).strip("_") or "model"
    return chosen, label


EXPORT_HEADER = [
    "Brand", "Model", "Variant", "Cashify Price", "DOFY Price", "Best Offer", "Difference", "Difference %",
    "Cashify Checked", "DOFY Checked", "Cashify URL", "DOFY URL",
]


def export_line(row: dict[str, Any], date_format: Any) -> list[Any]:
    cashify, dofy = row["cashify"], row["dofy"]
    return [
        row["brand"], row["model"], row["variant"], row["cashify_value"], row["dofy_value"], row["best"] or "",
        row["gap"], row["gap_pct"],
        date_format(cashify["date"]) if cashify else "", date_format(dofy["date"]) if dofy else "",
        cashify["url"] if cashify else "", dofy["url"] if dofy else "",
    ]


@app.get("/models.json")
def model_options_json() -> Any:
    """Model names for the export picker, fetched when the menu first opens so
    the list is not embedded in every page."""
    return jsonify(load_price_data()["model_options"])


@app.get("/export.csv")
def export_csv_now() -> Any:
    rows, label = export_rows(request.args.get("scope", "view"), request.args.get("model", ""))
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(EXPORT_HEADER)
    for row in rows:
        writer.writerow(export_line(row, lambda value: value or ""))
    return Response(
        buffer.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": f"attachment; filename=buyback_{label}.csv"},
    )


def excel_datetime(value: str | None) -> Any:
    """ISO timestamp -> naive local datetime, which Excel can sort and filter."""
    parsed = parse_iso(value)
    return parsed.astimezone(DISPLAY_TIMEZONE).replace(tzinfo=None) if parsed else (value or "")


def style_sheet(sheet: Any, widths: list[int], money_columns: list[int], date_columns: list[int]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill("solid", fgColor="1F3A8A")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for column in money_columns:
        for cell in sheet[get_column_letter(column)][1:]:
            cell.number_format = "#,##0"
    for column in date_columns:
        for cell in sheet[get_column_letter(column)][1:]:
            cell.number_format = "dd-mmm-yyyy hh:mm"
    sheet.freeze_panes = "A2"
    if sheet.max_row > 1:
        sheet.auto_filter.ref = sheet.dimensions


@app.get("/export.xlsx")
def export_excel_now() -> Any:
    from openpyxl import Workbook

    rows, label = export_rows(request.args.get("scope", "all"), request.args.get("model", ""))
    workbook = Workbook()

    compare = workbook.active
    compare.title = "Cashify vs DOFY"
    compare.append(EXPORT_HEADER)
    for row in rows:
        compare.append(export_line(row, excel_datetime))
    style_sheet(compare, [12, 34, 16, 14, 14, 12, 12, 13, 20, 20, 60, 60], money_columns=[4, 5, 7], date_columns=[9, 10])

    # One line per platform price, for filtering in Excel.
    prices = workbook.create_sheet("Prices")
    prices.append(["Platform", "Brand", "Model", "Variant", "Price", "Currency", "Last Checked", "Source URL"])
    for row in rows:
        for source in (row["cashify"], row["dofy"]):
            if source:
                prices.append([
                    source["platform"], row["brand"], source["model"], source["variant"], source["price_value"],
                    source.get("currency") or "INR", excel_datetime(source["date"]), source["url"],
                ])
    style_sheet(prices, [10, 12, 34, 16, 12, 10, 20, 60], money_columns=[5], date_columns=[7])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return Response(
        buffer.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=buyback_{label}.xlsx"},
    )


# --------------------------------------------------------------- scraping
@app.post("/check-now")
def check_now() -> Any:
    if HOSTED_READ_ONLY:
        return hosted_read_only_response()
    started = start_job("all", "all", "")
    if request.headers.get("Accept") == "application/json":
        return jsonify(current_job_state()), (202 if started else 409)
    return redirect(url_for("index"))


@app.post("/discover-scope")
def discover_scope_now() -> Any:
    if HOSTED_READ_ONLY:
        return hosted_read_only_response()
    platform = request.form.get("scrape_platform", "").strip().lower() or "all"
    brand = request.form.get("scrape_brand", "").strip() or "all"
    model = request.form.get("scrape_model", "").strip()[:100]
    platform = platform if platform in SCRAPE_PLATFORMS else "all"
    brand = brand if brand == "all" or brand in SCRAPE_BRANDS else "all"
    started = start_job(platform, brand, model)
    if request.headers.get("Accept") == "application/json":
        return jsonify(current_job_state()), (202 if started else 409)
    return redirect(url_for("index"))


@app.get("/job-status")
def job_status() -> Any:
    return jsonify(current_job_state())


def hosted_read_only_response() -> Any:
    payload = current_job_state()
    payload["message"] = "Hosted dashboard is read-only"
    payload["output"] = "Run the scraper on a PC or GitHub Actions; this host only shows the prices."
    if request.headers.get("Accept") == "application/json":
        return jsonify(payload), 409
    return redirect(url_for("index"))


def update_job_progress(line: str, output: str) -> None:
    """Fold one line of scraper output into the job state."""
    patch: dict[str, Any] = {}
    marker = line.find("PROGRESS ")
    if marker != -1:
        try:
            patch = json.loads(line[marker + len("PROGRESS "):])
        except ValueError:
            patch = {}
    with job_lock:
        progress = dict(job_state.get("progress") or {})
        for key in ("done", "rows", "errors", "current", "percent", "total"):
            if key in patch:
                progress[key] = patch[key]
        job_state["progress"] = progress
        job_state["output"] = output


def current_job_state() -> dict[str, Any]:
    with job_lock:
        return dict(job_state)


def job_scope(platform: str, brand: str, model: str) -> str:
    scope = SCRAPE_PLATFORMS.get(platform, platform)
    scope += " for all brands" if brand == "all" else f" for {brand}"
    if model:
        scope += f" matching {model}"
    return scope


def start_job(platform: str, brand: str, model: str) -> bool:
    with job_lock:
        if job_state["running"]:
            return False
        message = f"Checking {job_scope(platform, brand, model)}..."
        job_state.update({
            "id": str(uuid.uuid4()),
            "running": True,
            "started_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
            "finished_at": None,
            "returncode": None,
            "message": message,
            "output": "",
            "progress": {"done": 0, "rows": 0, "errors": 0, "current": "", "scope": message, "percent": 0},
        })
    thread = threading.Thread(target=run_scrape_command, args=(platform, brand, model), daemon=True)
    thread.start()
    return True


def scrape_command(platform: str, brand: str, model: str) -> list[str]:
    command = [
        sys.executable, str(ROOT / "buyback.py"), "scrape", "--platform", platform, "--brand", brand,
        "--delay", os.environ.get("SCRAPE_DELAY", "0.5"),
    ]
    if model:
        command.extend(["--model", model])
    if REMOTE_STORE is None:
        command.extend(["--db", str(db_path())])
    return command


def run_scrape_command(platform: str, brand: str, model: str) -> None:
    try:
        process = subprocess.Popen(
            scrape_command(platform, brand, model),
            cwd=ROOT,
            text=True,
            encoding="utf-8",
            errors="replace",
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            bufsize=1,
        )
        output_lines: list[str] = []
        assert process.stdout is not None
        for line in process.stdout:
            clean = line.strip()
            if not clean:
                continue
            output_lines.append(clean)
            output_lines = output_lines[-200:]
            update_job_progress(clean, "\n".join(output_lines)[-6000:])
        returncode = process.wait(timeout=6 * 3600)
        with job_lock:
            job_state.update({
                "running": False,
                "finished_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
                "returncode": returncode,
                "message": "Price check complete" if returncode == 0 else "Price check finished with errors",
            })
            job_state["progress"]["current"] = ""
    except Exception as exc:
        with job_lock:
            job_state.update({
                "running": False,
                "finished_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
                "returncode": -1,
                "message": f"Price check failed: {type(exc).__name__}",
                "output": str(exc),
            })


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=int(os.environ.get("PORT", "5050")), debug=True, use_reloader=False)
